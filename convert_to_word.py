"""
Convert Markdown to Word Document
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re


def create_word_document(md_file, output_file):
    """Convert markdown to Word document"""
    
    # Read markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Parse markdown line by line
    lines = content.split('\n')
    i = 0
    in_table = False
    table_data = []
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue
        
        # Heading 1
        if line.startswith('# ') and not line.startswith('##'):
            heading = line[2:].strip()
            p = doc.add_heading(heading, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Heading 2
        elif line.startswith('## '):
            heading = line[3:].strip()
            p = doc.add_heading(heading, level=2)
        
        # Heading 3
        elif line.startswith('### '):
            heading = line[4:].strip()
            p = doc.add_heading(heading, level=3)
        
        # Table detection
        elif '|' in line and not in_table:
            # Start of table
            in_table = True
            table_data = [line]
        
        elif in_table and '|' in line:
            # Continue table
            table_data.append(line)
        
        elif in_table and '|' not in line:
            # End of table - create it
            create_table(doc, table_data)
            in_table = False
            table_data = []
            continue
        
        # Code block
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            # Add code as paragraph with monospace font
            if code_lines:
                p = doc.add_paragraph('\n'.join(code_lines))
                p.style = 'Normal'
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
        
        # Bullet list
        elif line.startswith('- '):
            text = line[2:].strip()
            doc.add_paragraph(text, style='List Bullet')
        
        # Regular paragraph
        else:
            # Skip separator lines
            if not line.startswith('---') and not line.startswith('==='):
                doc.add_paragraph(line)
        
        i += 1
    
    # Handle any remaining table
    if in_table and table_data:
        create_table(doc, table_data)
    
    # Save document
    doc.save(output_file)
    print(f"✓ Word document created: {output_file}")


def create_table(doc, table_data):
    """Create a table from markdown data"""
    if not table_data or len(table_data) < 2:
        return
    
    # Parse table rows
    rows = []
    for line in table_data:
        if '---' in line or '===' in line:
            continue  # Skip separator
        cells = [cell.strip() for cell in line.split('|')]
        cells = [c for c in cells if c]  # Remove empty
        if cells:
            rows.append(cells)
    
    if not rows:
        return
    
    # Determine number of columns
    max_cols = max(len(row) for row in rows)
    
    # Create Word table
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = 'Light Grid Accent 1'
    
    # Fill table
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < max_cols:
                cell = row.cells[j]
                cell.text = cell_text
                
                # Bold header row
                if i == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
    
    # Add spacing after table
    doc.add_paragraph()


if __name__ == "__main__":
    create_word_document(
        'SYSTEM_ARCHITECTURE_DETAILED.md',
        'SYSTEM_ARCHITECTURE_DETAILED.docx'
    )
