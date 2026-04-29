"""
Generate test cases for a specific module + focus area
Called by run_all_focus_areas.py with module_id and focus_area arguments

Usage:
  python generate_all_login_tests.py <focus_area> <module_id>
  e.g. python generate_all_login_tests.py authentication login
  e.g. python generate_all_login_tests.py filters hcp_search
"""
import ast
import os
import sys
import json
from ai_agent.agents.planner import PlannerAgent
from ai_agent.agents.designer import DesignerAgent

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEST_FILE  = os.path.join(BASE_DIR, "tests", "test_login.py")
EXCEL_FILE = os.path.join(BASE_DIR, "Testcases", "all_login_test_cases_COMPLETE.xlsx")

CLASS_LABELS = {
    "TestLoginPositive":  "Positive",
    "TestLoginEdgeCases": "Edge Case",
    "TestLoginNegative":  "Negative / Security",
    "TestLoginSession":   "Session / Post-Login",
}


RESULTS_FILE = os.path.join(BASE_DIR, ".pytest_results.txt")  # written by conftest after each run

# Hand-authored Expected Result per test — matches the assert logic in test_login.py
EXPECTED_RESULTS = {
    "test_valid_login_redirects_to_dashboard":      "User is redirected to dashboard; URL contains trinitylifesciences.com",
    "test_valid_login_email_case_insensitive":       "SSO accepts uppercase email and advances to password step",
    "test_valid_login_email_with_spaces_trimmed":    "Login succeeds; spaces trimmed from email automatically",
    "test_login_page_loads_correctly":               "Sign In button is visible on the app login page",
    "test_forgot_password_link_visible":             "'Can't access your account?' link is visible on SSO password page",
    "test_forgot_password_link_navigates":           "Clicking recovery link navigates to microsoftonline password reset page",
    "test_password_field_is_masked":                 "Password input type=password (field is masked)",
    "test_login_with_empty_email":                   "User is NOT logged in; SSO shows validation error",
    "test_login_with_empty_password":                "User is NOT logged in; SSO shows validation error",
    "test_login_email_with_plus_sign":               "SSO stays on microsoftonline; email format with + is accepted",
    "test_login_tab_navigation":                     "Focus moves to INPUT, BUTTON, or A element after Tab key",
    "test_login_enter_key_submits":                  "Enter key advances to password step on SSO",
    "test_login_page_title":                         "Page title is not empty",
    "test_browser_back_after_login":                 "Browser back stays on trinitylifesciences.com or microsoftonline.com",
    "test_invalid_password_shows_error":             "User is NOT logged in; SSO shows incorrect password error",
    "test_invalid_email_shows_error":                "User is NOT logged in; SSO shows account not found error",
    "test_invalid_email_format":                     "User is NOT logged in; SSO shows email format error",
    "test_sql_injection_in_email":                   "User is NOT logged in; SQL injection blocked by SSO",
    "test_sql_injection_in_password":                "User is NOT logged in; SQL injection in password blocked",
    "test_xss_in_email_field":                       "User is NOT logged in; XSS script not executed",
    "test_very_long_email_rejected":                 "User is NOT logged in; 500+ char email rejected by SSO",
    "test_special_characters_in_password":           "User is NOT logged in; wrong special-char password rejected",
    "test_remember_me_session_persists":             "User is logged in after clicking Yes on 'Stay signed in?' prompt",
    "test_login_access_hcp_targeting_module":        "Dashboard loads with HCP/Specialty/Segmentation content visible",
    "test_mfa_prompt_appears_after_password":        "MFA prompt OR 'Stay signed in?' OR successful redirect appears",
    "test_session_expiry_redirects_to_login":        "Unauthenticated access redirects to SSO or shows Sign In button",
    "test_logout_clears_session":                    "After logout, protected route redirects to SSO or /login",
    "test_account_lockout_after_multiple_failures":  "User is NOT logged in after 3 wrong password attempts",
    "test_concurrent_session_same_browser":          "Second tab lands on app or SSO, not an error page",
}


def _load_last_results() -> dict:
    """Read {func_name: PASS|FAIL|SKIP} written by conftest after last pytest run."""
    results = {}
    if not os.path.exists(RESULTS_FILE):
        return results
    with open(RESULTS_FILE) as f:
        for line in f:
            parts = line.strip().split("|", 1)
            if len(parts) == 2:
                results[parts[0]] = parts[1]
    return results


def rebuild_excel_from_tests():
    """
    Rebuild Testcases/all_login_test_cases_COMPLETE.xlsx directly from
    tests/test_login.py — includes Expected Result + Actual Result from last run.
    Run: python test_generation_scripts/generate_all_login_tests.py --rebuild
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    # Parse test functions via AST
    with open(TEST_FILE, "r", encoding="utf-8") as f:
        tree = ast.parse(f.read())

    tests = []
    for node in ast.walk(tree):
        if isinstance(node, ast.ClassDef):
            cls = node.name
            for item in node.body:
                if isinstance(item, ast.FunctionDef) and item.name.startswith("test_"):
                    tests.append({
                        "fn":       item.name,
                        "category": CLASS_LABELS.get(cls, cls),
                        "doc":      (ast.get_docstring(item) or "").strip(),
                    })

    last_results = _load_last_results()
    run_ts = last_results.get("__timestamp__", "Not run yet")

    STATUS_COLOUR = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "SKIP": "FFEB9C"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Login Test Cases"

    headers = ["#", "Test ID", "Category", "Description",
               "Expected Result", "Actual Result", "Status", "Last Run"]
    ws.append(headers)

    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = hdr_fill
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, t in enumerate(tests, 1):
        fn      = t["fn"]
        status  = last_results.get(fn, "")
        exp_res = EXPECTED_RESULTS.get(fn, t["doc"])
        act_res = status   # PASS / FAIL / SKIP — will be enriched by future runs

        ws.append([i, fn, t["category"], t["doc"],
                   exp_res, act_res, status, run_ts if status else ""])

        # Colour Status and Actual Result cells
        if status in STATUS_COLOUR:
            col = STATUS_COLOUR[status]
            fill = PatternFill(start_color=col, end_color=col, fill_type="solid")
            ws.cell(ws.max_row, 6).fill = fill
            ws.cell(ws.max_row, 7).fill = fill
            ws.cell(ws.max_row, 6).font = Font(bold=True)
            ws.cell(ws.max_row, 7).font = Font(bold=True)

        for c in range(1, len(headers) + 1):
            ws.cell(ws.max_row, c).alignment = Alignment(vertical="top", wrap_text=True)

    for i, w in enumerate([4, 45, 22, 50, 55, 12, 10, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(EXCEL_FILE), exist_ok=True)
    wb.save(EXCEL_FILE)

    passed  = sum(1 for v in last_results.values() if v == "PASS")
    failed  = sum(1 for v in last_results.values() if v == "FAIL")
    skipped = sum(1 for v in last_results.values() if v == "SKIP")
    print(f"Updated  → {EXCEL_FILE}")
    print(f"Tests    : {len(tests)} rows")
    print(f"Last run : {run_ts}  |  PASS={passed}  FAIL={failed}  SKIP={skipped}")

CONFIG_FILE = "test_generation_scripts/modules_config.json"

def load_module_config(module_id):
    """Load UI data for a specific module from config"""
    with open(CONFIG_FILE, 'r') as f:
        config = json.load(f)
    for module in config['modules']:
        if module['module_id'] == module_id:
            return module, config.get('focus_area_instructions', {})
    return None, {}

def main(focus_area=None, module_id="login"):
    # Load UI data from config
    module_config, focus_instructions = load_module_config(module_id)
    
    if module_config:
        ui_data = module_config['ui_data'].copy()
        module_name = module_config['name']
    else:
        # Fallback: default login UI data
        module_name = "Login"
        ui_data = {
            "url": "https://ce-ts-dev.trinitylifesciences.com/login",
            "module": "Login",
            "inputs": [
                {"label": "Email", "type": "email", "placeholder": "Enter email"},
                {"label": "Password", "type": "password", "placeholder": "Enter password"}
            ],
            "buttons": [{"text": "Login", "type": "submit"}],
            "links": [{"text": "Forgot Password?", "href": "/forgot-password"}]
        }

    # Inject focus area into UI data for planner
    ui_data['focus_area'] = focus_area
    if focus_area and focus_area in focus_instructions:
        ui_data['focus_instruction'] = focus_instructions[focus_area]
    
    print("="*80)
    print(f"GENERATING: {module_name} | Focus: {focus_area or 'general'}")
    print("="*80)
    
    # Step 1: Generate scenarios
    print("\n[1/3] Generating test scenarios...")
    planner = PlannerAgent()
    scenarios = planner.generate_scenarios(ui_data)
    
    all_scenarios = []
    all_scenarios.extend([(s, "positive") for s in scenarios.get("positive_scenarios", [])])
    all_scenarios.extend([(s, "edge_case") for s in scenarios.get("edge_cases", [])])
    all_scenarios.extend([(s, "negative") for s in scenarios.get("negative_scenarios", [])])
    
    print(f"   [OK] Generated {len(all_scenarios)} scenarios")
    
    # Step 2: Convert each scenario to test case
    print("\n[2/3] Converting scenarios to test cases...")
    designer = DesignerAgent()
    all_test_cases = []
    
    for idx, (scenario, scenario_type) in enumerate(all_scenarios, 1):
        print(f"   Processing {idx}/{len(all_scenarios)}: {scenario_type}")
        single_scenario = {scenario_type + "_scenarios": [scenario]}
        try:
            test_cases = designer.design_tests(single_scenario, ui_data, max_tests=1)
            if test_cases:
                test_case = test_cases[0]
                test_case['scenario_type'] = scenario_type
                test_case['scenario_text'] = scenario
                all_test_cases.append(test_case)
                print(f"      [OK] Created: {test_case.get('name', 'Unknown')}")
            else:
                print(f"      [WARN] No test case generated")
        except Exception as e:
            print(f"      [ERROR] {str(e)[:100]}")
    
    print(f"\n   [OK] Generated {len(all_test_cases)} test cases")
    
    # Step 3: Save JSON (Excel merged later by run_all_focus_areas.py)
    print("\n[3/3] Saving results...")
    os.makedirs("outputs", exist_ok=True)
    
    json_file = f"outputs/{module_id}_{focus_area}.json" if focus_area else f"outputs/{module_id}_general.json"
    with open(json_file, 'w') as f:
        json.dump({
            "module": module_name,
            "focus_area": focus_area or "general",
            "scenarios": scenarios,
            "test_cases": all_test_cases,
            "summary": {
                "total_scenarios": len(all_scenarios),
                "total_test_cases": len(all_test_cases),
                "positive": len([tc for tc in all_test_cases if tc.get('scenario_type') == 'positive']),
                "edge_cases": len([tc for tc in all_test_cases if tc.get('scenario_type') == 'edge_case']),
                "negative": len([tc for tc in all_test_cases if tc.get('scenario_type') == 'negative'])
            }
        }, f, indent=2)
    
    print(f"   [OK] Saved {len(all_test_cases)} test cases to {json_file}")

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--rebuild":
        rebuild_excel_from_tests()
    else:
        focus_area = sys.argv[1] if len(sys.argv) > 1 else None
        module_id  = sys.argv[2] if len(sys.argv) > 2 else "login"
        main(focus_area, module_id)
